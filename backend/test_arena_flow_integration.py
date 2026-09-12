import io
from fastapi.testclient import TestClient
from docx import Document
from backend.main import app

client = TestClient(app)

def create_sample_docx(topic: str, text: str) -> bytes:
    doc = Document()
    doc.add_heading(f"Advanced Study: {topic}", 0)
    doc.add_paragraph(text)
    doc.add_paragraph("Key invariant: In asynchronous state management, state mutations are scheduled and reconciled in batches.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def test_checkpoint3_end_to_end():
    print("\n--- CHECKPOINT 3: FULL END-TO-END FLOW VERIFICATION ---")

    # Flow 1: Topic Path -> Quiz -> Autopsy
    print("\n[FLOW 1] Testing Topic Path -> Quiz -> Autopsy...")
    res_topic = client.post("/api/generate", json={"topic": "React State & Closures", "num_questions": 3})
    assert res_topic.status_code == 200
    topic_quiz = res_topic.json()
    assert len(topic_quiz["questions"]) >= 3
    q1 = topic_quiz["questions"][0]
    print(f"  -> Generated topic quiz '{topic_quiz['title']}' with {len(topic_quiz['questions'])} questions.")
    
    # Trigger Autopsy on an incorrect option
    wrong_opt = next(o for o in q1["options"] if not o["is_correct"])
    correct_opt = next(o for o in q1["options"] if o["is_correct"])
    
    res_autopsy = client.post("/api/autopsy", json={
        "question_id": q1["id"],
        "selected_option_id": wrong_opt["id"],
        "question_prompt": q1["prompt"],
        "chosen_text": wrong_opt["text"],
        "correct_text": correct_opt["text"],
        "cognitive_trap_name": q1["cognitive_trap_name"]
    })
    assert res_autopsy.status_code == 200
    autopsy_data = res_autopsy.json()
    assert "fallacy_name" in autopsy_data
    assert "ten_second_cure" in autopsy_data
    print(f"  -> Topic autopsy diagnosed: '{autopsy_data['fallacy_name']}'")
    print(f"     10-Sec Cure: '{autopsy_data['ten_second_cure']}'")

    # Flow 2: Document Path -> Quiz -> Autopsy
    print("\n[FLOW 2] Testing Document Path -> Quiz -> Autopsy...")
    docx_bytes = create_sample_docx(
        "Distributed Consensus",
        "Paxos and Raft achieve consensus in asynchronous networks with fail-stop failures using majority quorums."
    )
    files = {"file": ("distributed_consensus_notes.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res_doc = client.post("/quiz/generate-from-document", files=files, data={"num_questions": "3"})
    assert res_doc.status_code == 200
    doc_quiz = res_doc.json()
    assert len(doc_quiz["questions"]) >= 1
    assert doc_quiz["source_document"] == "distributed_consensus_notes.docx"
    doc_q1 = doc_quiz["questions"][0]
    print(f"  -> Generated document quiz with {len(doc_quiz['questions'])} questions from '{doc_quiz['source_document']}'.")
    print(f"     Summary: '{doc_quiz['source_summary']}'")

    # Trigger Autopsy on document question
    doc_wrong_opt = next(o for o in doc_q1["options"] if not o["is_correct"])
    doc_correct_opt = next(o for o in doc_q1["options"] if o["is_correct"])

    res_doc_autopsy = client.post("/api/autopsy", json={
        "question_id": doc_q1["id"],
        "selected_option_id": doc_wrong_opt["id"],
        "question_prompt": doc_q1["prompt"],
        "chosen_text": doc_wrong_opt["text"],
        "correct_text": doc_correct_opt["text"],
        "cognitive_trap_name": doc_q1["cognitive_trap_name"],
        "source_document": doc_q1.get("source_document")
    })
    assert res_doc_autopsy.status_code == 200
    doc_autopsy_data = res_doc_autopsy.json()
    assert "fallacy_name" in doc_autopsy_data
    assert "ten_second_cure" in doc_autopsy_data
    print(f"  -> Document autopsy diagnosed: '{doc_autopsy_data['fallacy_name']}'")
    print(f"     10-Sec Cure: '{doc_autopsy_data['ten_second_cure']}'")

    print("\nALL CHECKPOINT 3 END-TO-END FLOWS VERIFIED SUCCESSFULLY!")

if __name__ == "__main__":
    test_checkpoint3_end_to_end()
