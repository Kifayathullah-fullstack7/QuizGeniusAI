import io
from fastapi.testclient import TestClient
from docx import Document
from backend.main import app

client = TestClient(app)

def create_sample_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 712 Td ({text}) Tj ET".encode("latin-1", "replace")
    pdf_template = (
        b"%PDF-1.4\n"
        b"1 0 obj << /Type /Catalog /Pages 2 0 R >> endobj\n"
        b"2 0 obj << /Type /Pages /Kids [3 0 R] /Count 1 >> endobj\n"
        b"3 0 obj << /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >> endobj\n"
        b"4 0 obj << /Length " + str(len(stream)).encode() + b" >> stream\n"
        + stream + b"\nendstream\nendobj\n"
        b"5 0 obj << /Type /Font /Subtype /Type1 /BaseFont /Helvetica >> endobj\n"
        b"xref\n0 6\n0000000000 65535 f \n0000000009 00000 n \n0000000058 00000 n \n0000000115 00000 n \n0000000244 00000 n \n0000000330 00000 n \n"
        b"trailer << /Size 6 /Root 1 0 R >>\nstartxref\n410\n%%EOF"
    )
    return pdf_template

def create_sample_docx(text: str) -> bytes:
    doc = Document()
    doc.add_heading("PostgreSQL Isolation Levels and Concurrency", 0)
    doc.add_paragraph(text)
    doc.add_paragraph("MVCC creates tuple versions to prevent readers from blocking writers in high-throughput databases.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()

def run_tests():
    print("\n--- CHECKPOINT 2: POST /quiz/generate-from-document TESTS ---")

    # 1. Test with real PDF
    print("[TEST 1] Testing POST /quiz/generate-from-document with real PDF...")
    pdf_bytes = create_sample_pdf("Distributed hash tables utilize consistent hashing to minimize remapping when nodes fail.")
    files = {"file": ("distributed_hash_tables.pdf", pdf_bytes, "application/pdf")}
    data = {"num_questions": "3", "difficulty": "hard"}
    res = client.post("/quiz/generate-from-document", files=files, data=data)
    assert res.status_code == 200, f"Expected 200, got {res.status_code}: {res.text}"
    json_data = res.json()
    assert "questions" in json_data, "Missing 'questions' in response"
    assert len(json_data["questions"]) >= 1, "Expected at least 1 question"
    assert "source_summary" in json_data, "Missing 'source_summary' in response"
    assert json_data["source_document"] == "distributed_hash_tables.pdf"
    assert json_data["is_grounded"] is True
    print(f"  -> Passed! Generated {len(json_data['questions'])} questions. Summary: {json_data['source_summary']}")

    # 2. Test with real DOCX
    print("[TEST 2] Testing POST /quiz/generate-from-document with real DOCX...")
    docx_bytes = create_sample_docx("Repeatable Read isolation prevents phantom reads in SQL databases using snapshot isolation.")
    files = {"file": ("database_isolation.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}
    res = client.post("/quiz/generate-from-document", files=files, data={"num_questions": "4"})
    assert res.status_code == 200, f"Expected 200, got {res.status_code}: {res.text}"
    json_data = res.json()
    assert len(json_data["questions"]) >= 1
    assert "source_summary" in json_data
    print(f"  -> Passed! DOCX Quiz generated. Questions: {len(json_data['questions'])}")

    # 3. Test with unsupported format
    print("[TEST 3] Testing POST /quiz/generate-from-document with unsupported format (.png)...")
    files = {"file": ("diagram.png", b"\x89PNG\r\n\x1a\nfakeimagecontent", "image/png")}
    res = client.post("/quiz/generate-from-document", files=files)
    assert res.status_code == 400, f"Expected 400, got {res.status_code}"
    print(f"  -> Passed! Rejected unsupported file with 400: {res.json().get('detail')}")

    # 4. Test with garbage / empty document
    print("[TEST 4] Testing POST /quiz/generate-from-document with empty text document...")
    files = {"file": ("empty.txt", b"", "text/plain")}
    res = client.post("/quiz/generate-from-document", files=files)
    assert res.status_code == 400, f"Expected 400, got {res.status_code}"
    print(f"  -> Passed! Rejected empty document with 400: {res.json().get('detail')}")

    # 5. Verify existing topic-based route
    print("[TEST 5] Testing existing POST /quiz/generate topic path for zero regression...")
    res = client.post("/quiz/generate", json={"topic": "React State & Closures", "num_questions": 3})
    assert res.status_code == 200, f"Expected 200, got {res.status_code}"
    assert len(res.json()["questions"]) >= 3
    print("  -> Passed! Topic path works as expected.")

    print("\nALL CHECKPOINT 2 TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    run_tests()
